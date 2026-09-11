import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_to_excel(cards_data: list, output_excel_path: str):
    """
    將結構化名片清單輸出為精美排版的商務 Excel 活頁簿
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "名片清冊"

    headers = [
        "編號", "公司名稱", "姓名", "職稱", "行動電話", 
        "公司電話", "傳真", "Email", "地址", "統一編號", 
        "官方網站", "來源頁碼", "名片圖檔", "原始辨識文字"
    ]
    ws.append(headers)

    # 專業商務風格表頭 (深藍色底白字)
    header_font = Font(name="微軟正黑體", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    data_font = Font(name="微軟正黑體", size=10)
    alt_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    for row_idx, data in enumerate(cards_data, start=2):
        row_vals = [
            data.get("編號", row_idx - 1),
            data.get("公司名稱", ""),
            data.get("姓名", ""),
            data.get("職稱", ""),
            data.get("行動電話", ""),
            data.get("公司電話", ""),
            data.get("傳真", ""),
            data.get("Email", ""),
            data.get("地址", ""),
            data.get("統一編號", ""),
            data.get("網站", ""),
            data.get("來源頁碼", ""),
            data.get("圖檔檔名", ""),
            data.get("原始文字", "")
        ]
        ws.append(row_vals)
        fill = alt_fill if row_idx % 2 == 0 else white_fill
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.fill = fill
            cell.border = thin_border
            if col_idx in [1, 5, 6, 7, 10, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # 自動調整欄寬
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            c_len = sum(2 if ord(char) > 127 else 1 for char in val_str)
            if c_len > max_len:
                max_len = c_len
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    ws.row_dimensions[1].height = 26
    for r in range(2, len(cards_data) + 2):
        ws.row_dimensions[r].height = 22

    wb.save(output_excel_path)
